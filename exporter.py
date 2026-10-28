"""Write result tables to an .xlsx workbook (one sheet per uploaded file)."""
import re
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

import settings as cfg

_FONT = "Arial"
_HEADER_FONT = Font(name=_FONT, bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", start_color="FF2B2F45")
_RED_FILL = PatternFill("solid", start_color="FFD32F2F")
_RED_FONT = Font(name=_FONT, bold=True, color="FFFFFF")
_BODY_FONT = Font(name=_FONT)
_TOP_WRAP = Alignment(vertical="top", wrap_text=True)

_WIDTHS = {
    cfg.COL_NUMBER: 12, cfg.COL_ID: 30, cfg.COL_SOURCE: 60, cfg.COL_TARGET: 60,
    cfg.COL_BOTH_EN: 22, cfg.COL_EN_MATCH: 20, cfg.COL_SRC_LANG: 18, cfg.COL_TGT_LANG: 18,
    cfg.COL_CONF: 14,
}
_BAD_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


def _sheet_title(filename: str, used: set[str]) -> str:
    base = _BAD_SHEET_CHARS.sub("_", Path(filename).stem).strip("' ")[:31] or "Sheet"
    title, n = base, 2
    while title.lower() in used:
        suffix = f"_{n}"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


def _write_sheet(ws, df: pd.DataFrame) -> None:
    for col, name in enumerate(cfg.RESULT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = _WIDTHS[name]
    ws.row_dimensions[1].height = 45

    for row, values in enumerate(df.itertuples(index=False, name=None), start=2):
        for col, (name, value) in enumerate(zip(cfg.RESULT_COLUMNS, values), start=1):
            cell = ws.cell(row=row, column=col)
            cell.font, cell.alignment = _BODY_FONT, _TOP_WRAP
            if name == cfg.COL_NUMBER:
                cell.value = int(value)
            elif name == cfg.COL_CONF:
                if pd.notna(value):
                    cell.value = float(value)
                    cell.number_format = "0.0%"
                    if value < cfg.CONFIDENCE_THRESHOLD:
                        cell.fill, cell.font = _RED_FILL, _RED_FONT
            else:
                # Force plain text so segments starting with "=" are never run as formulas.
                cell.value = ILLEGAL_CHARACTERS_RE.sub("", str(value))
                cell.data_type = "s"
                if name == cfg.COL_EN_MATCH and value == "English not matching":
                    cell.fill, cell.font = _RED_FILL, _RED_FONT

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build_workbook(sheets: list[tuple[str, pd.DataFrame]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for filename, df in sheets:
        _write_sheet(wb.create_sheet(_sheet_title(filename, used)), df)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
